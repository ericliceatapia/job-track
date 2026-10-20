from datetime import date

import pytest
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from jobtrack.models import Base, Application, ApplicationEvent
from jobtrack import services as s
from jobtrack.exporting import export_applications

DAY = date(2026, 10, 1)

@pytest.fixture
def session():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session
    engine.dispose()

@pytest.fixture
def make(session):
    def create(**changes):
        fields = dict(company=" Acme ", role_title=" Analyst ", category="analytics",
                      location="  ", work_modality="remote", job_type="full-time",
                      source="linkedin", application_date=DAY, job_id="00123")
        fields.update(changes)
        return s.create_application(session, **fields)
    return create

def test_creation(session, make):
    a = make()
    assert (a.company, a.role_title, a.category, a.location, a.status) == (
        "Acme", "Analyst", "Analytics", None, "Applied")
    assert [(e.event_type, e.event_date) for e in a.events] == [("APPLICATION_CREATED", DAY)]

@pytest.mark.parametrize("status", ["Offer", "Rejected", "Withdrawn"])
def test_terminal_status(session, make, status):
    a = make()
    s.change_status(session, a.id, f" {status.lower()} ", DAY)
    assert a.status == status
    for action in [lambda: s.change_status(session, a.id, "Applied", DAY),
                   lambda: s.change_status(session, a.id, status, DAY),
                   lambda: s.record_oa(session, a.id, "received", DAY),
                   lambda: s.record_interview(session, a.id, DAY)]:
        with pytest.raises(ValueError):
            action()
    s.edit_role_notes(session, a.id, " role ")
    s.record_note(session, a.id, " activity ", DAY)
    assert a.notes == "role"

def test_oa_and_interviews(session, make):
    a = make()
    with pytest.raises(ValueError):
        s.record_oa(session, a.id, "completed", DAY)
    for action in ("received", "completed"):
        assert s.record_oa(session, a.id, f" {action.upper()} ", DAY).event_type == f"OA_{action.upper()}"
        with pytest.raises(ValueError):
            s.record_oa(session, a.id, action, DAY)
    assert s.record_interview(session, a.id, DAY).id != s.record_interview(session, a.id, DAY).id

def test_notes_and_missing(session, make):
    a = make()
    s.edit_role_notes(session, a.id, " role ")
    s.edit_role_notes(session, a.id, None)
    assert a.notes is None
    assert s.record_note(session, a.id, " text ", DAY).notes == "text"
    with pytest.raises(ValueError):
        s.record_note(session, a.id, " ", DAY)
    assert s.get_application(session, 999) is None
    with pytest.raises(ValueError):
        s.edit_role_notes(session, 999, "x")

def test_filters(session, make):
    a = make(company="100% Acme")
    make(company="Other", category="software")
    assert s.list_applications(session, company="%", category=" ANALYTICS ", status="applied") == [a]
    assert s.list_applications(session, company="acme", category="software") == []
    assert [x.id for x in s.list_applications(session)] == [1, 2]
    with pytest.raises(ValueError):
        s.list_applications(session, status="invalid")

def test_export(session, make, tmp_path):
    a = make(company="=literal")
    s.record_note(session, a.id, "=note", date(2026, 9, 1))
    path = tmp_path / "result.xlsx"
    assert export_applications(s.list_applications(session), path) == 2
    book = load_workbook(path)
    assert book.sheetnames == ["Applications", "History"]
    assert book["Applications"]["B2"].data_type == "s"
    assert book["Applications"]["K2"].value == "00123"
    assert book["History"]["F2"].value == "NOTE_ADDED"
    book.close()
    with pytest.raises(ValueError):
        export_applications([a], path)
    assert export_applications([], path, overwrite=True) == 0

def test_failed_write_rolls_back(session, make):
    def fail(mapper, connection, target):
        raise RuntimeError("injected event failure")
    event.listen(ApplicationEvent, "before_insert", fail)
    try:
        with pytest.raises(RuntimeError):
            make()
    finally:
        event.remove(ApplicationEvent, "before_insert", fail)
    assert session.scalars(select(Application)).all() == []
    assert session.scalars(select(ApplicationEvent)).all() == []
    assert make().id == 1
