from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from jobtrack.models import Application
from jobtrack.errors import ValidationError


def export_applications(
    applications: list[Application],
    destination: Path,
    *,
    overwrite: bool = False,
) -> int:
    """Write an Excel workbook and return the number of exported events."""

    if destination.suffix.lower() != ".xlsx":
        raise ValidationError("Export filename must end in .xlsx.")
    if destination.exists() and not overwrite:
        raise ValidationError("File already exists. Enable overwrite to replace it.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append([
        "Application ID", "Company", "Role Title", "Category", "Location",
        "Work Modality", "Job Type", "Source", "Application Date", "Status",
        "Job ID", "Role Notes", "Created At", "Updated At",
    ])
    history = workbook.create_sheet("History")
    history.append([
        "Event ID", "Application ID", "Company", "Role Title", "Event Date",
        "Event Type", "Old Value", "New Value", "Activity Note", "Created At",
    ])
    events = []
    for application in applications:
        sheet.append([
            application.id, application.company, application.role_title,
            application.category, application.location, application.work_modality,
            application.job_type, application.source, application.application_date,
            application.status, application.job_id, application.notes,
            application.created_at, application.updated_at,
        ])
        events.extend((event, application) for event in application.events)
    for event, application in sorted(events, key=lambda pair: (pair[0].event_date, pair[0].id)):
        history.append([
            event.id, application.id, application.company, application.role_title,
            event.event_date, event.event_type, event.old_value, event.new_value,
            event.notes, event.created_at,
        ])

    for worksheet in workbook:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for row in worksheet:
            for cell in row:
                if isinstance(cell.value, str):
                    # User text must remain literal, including leading '='.
                    cell.data_type = "s"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24475B")
        for column in worksheet.columns:
            longest = 0
            for cell in column:
                value = cell.value
                if isinstance(value, datetime):
                    text = value.strftime("%Y-%m-%d %H:%M:%S")
                elif isinstance(value, date):
                    text = value.isoformat()
                else:
                    text = str(value) if value is not None else ""
                longest = max(longest, max(map(len, text.splitlines()), default=0))
            width = min(60, max(10, longest + 3))
            worksheet.column_dimensions[get_column_letter(column[0].column)].width = width
    for row in sheet.iter_rows(min_row=2, min_col=11, max_col=11):
        row[0].number_format = "@"

    payload = BytesIO()
    workbook.save(payload)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("wb" if overwrite else "xb") as output:
        output.write(payload.getvalue())
    return len(events)
